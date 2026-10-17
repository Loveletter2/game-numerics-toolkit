# -*- coding: utf-8 -*-
"""把「心 / 锁暝 / 守岸人」的输出轴表生成为一个自包含的 HTML 案例页。

设计原则：
1. 单文件、零依赖 —— 打开就能看，可以直接挂 GitHub Pages
2. 图表全部是内联 SVG —— 不依赖 JS 图表库，不会因为 CDN 挂掉变成空白
3. 数据在生成时算好写死 —— 页面打开不需要请求任何接口

用法:
    python case_study/build_demo.py <xlsx> <out.html>
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

# 与手拉表里「总伤」公式一致的归属划分
GROUPS = {
    "守岸人": {"rows": [23], "column": "E", "color": "#2e7d32"},
    "锁暝": {"rows": [24, 25, 31, 32, 33, 34, 39], "column": "E", "color": "#1f4e79"},
    "心": {"rows": [26, 27, 28, 29, 30, 35, 36, 37, 38], "column": "G", "color": "#c05000"},
}
FIGHT_SECONDS = 25

W, H = 860, 300
PAD_L, PAD_R, PAD_T, PAD_B = 268, 96, 24, 16


def load(xlsx: Path):
    wb = load_workbook(xlsx, data_only=True)
    ws = wb.active
    rows = {}
    for r in range(23, 42):
        label = ws.cell(row=r, column=1).value
        dtype = ws.cell(row=r, column=4).value
        expected = ws.cell(row=r, column=5).value
        crit = ws.cell(row=r, column=7).value
        rows[r] = {"label": label or f"row {r}", "type": dtype,
                   "expected": expected, "crit": crit}
    totals = {}
    for r in range(46, 51):
        name = ws.cell(row=r, column=1).value
        val = ws.cell(row=r, column=4).value
        if name:
            totals[str(name)] = val
    return rows, totals


def build(rows, totals):
    actions = []
    for owner, spec in GROUPS.items():
        for r in spec["rows"]:
            row = rows.get(r)
            if not row:
                continue
            value = row["expected"] if spec["column"] == "E" else row["crit"]
            if isinstance(value, (int, float)):
                actions.append({"owner": owner, "row": r, "label": row["label"],
                                "type": row["type"] or "-", "damage": float(value),
                                "column": spec["column"],
                                "color": GROUPS[owner]["color"]})
    actions.sort(key=lambda a: -a["damage"])

    per_owner = {k: 0.0 for k in GROUPS}
    for a in actions:
        per_owner[a["owner"]] += a["damage"]
    total = sum(per_owner.values())

    # 与手拉表自带的总伤公式交叉验证
    check = []
    for owner, key in (("守岸人", "守岸人总伤"), ("锁暝", "锁暝总伤"), ("心", "心总伤")):
        mine = per_owner[owner]
        theirs = totals.get(key)
        if isinstance(theirs, (int, float)) and theirs:
            check.append({"owner": owner, "mine": mine, "sheet": float(theirs),
                          "diff": (mine / float(theirs) - 1) * 100})
    return actions, per_owner, total, check


def fmt(n: float) -> str:
    if n >= 1e8:
        return f"{n/1e8:.2f}亿"
    if n >= 1e4:
        return f"{n/1e4:.1f}万"
    return f"{n:,.0f}"


def bar_svg(actions, total):
    seg_h = (H - PAD_T - PAD_B) / max(len(actions), 1)
    seg_h = min(seg_h, 26)
    height = PAD_T + PAD_B + seg_h * len(actions)
    inner = W - PAD_L - PAD_R
    out = [f'<svg viewBox="0 0 {W} {height}" width="100%" role="img" '
           f'aria-label="每个动作的伤害贡献">']
    for i, a in enumerate(actions):
        y = PAD_T + i * seg_h
        w = a["damage"] / actions[0]["damage"] * inner
        out.append(f'<rect x="{PAD_L}" y="{y+2:.1f}" width="{w:.1f}" '
                   f'height="{seg_h-4:.1f}" rx="3" fill="{a["color"]}" opacity="0.85"/>')
        label = a["label"][:18]
        out.append(f'<text x="{PAD_L-10}" y="{y+seg_h/2+4:.1f}" text-anchor="end" '
                   f'font-size="12" fill="#333">{label}</text>')
        pct = a["damage"] / total * 100
        out.append(f'<text x="{PAD_L+w+8:.1f}" y="{y+seg_h/2+4:.1f}" '
                   f'font-size="11" fill="#666">{fmt(a["damage"])} · {pct:.1f}%</text>')
    out.append("</svg>")
    return "\n".join(out)


def owner_svg(per_owner, total):
    w, h = W, 112
    x = 0
    out = [f'<svg viewBox="0 0 {w} {h}" width="100%" role="img" '
           f'aria-label="三人伤害占比">']
    for owner, value in sorted(per_owner.items(), key=lambda kv: -kv[1]):
        seg = value / total * w
        out.append(f'<rect x="{x:.1f}" y="10" width="{seg:.1f}" height="34" '
                   f'fill="{GROUPS[owner]["color"]}" opacity="0.85"/>')
        if seg > 90:
            out.append(f'<text x="{x+seg/2:.1f}" y="32" text-anchor="middle" '
                       f'font-size="13" fill="#fff">{owner} {value/total*100:.0f}%</text>')
        x += seg
    out.append('<text x="0" y="66" font-size="11" fill="#888">%s</text>'
               % (u'总伤 %s，分摊在 %d 秒轴上' % (fmt(total), FIGHT_SECONDS)))
    lx = 0.0
    for owner, value in sorted(per_owner.items(), key=lambda kv: -kv[1]):
        color = GROUPS[owner]["color"]
        out.append('<circle cx="%.1f" cy="96" r="5" fill="%s"/>' % (lx + 5, color))
        out.append('<text x="%.1f" y="100" font-size="12" fill="#333">%s %s\uff08%.1f%%\uff09</text>'
                   % (lx + 15, owner, fmt(value), value / total * 100))
        lx += 165
    out.append("</svg>")
    return "\n".join(out)


HTML = """<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>输出轴案例：心 / 锁暝 / 守岸人</title>
<style>
  :root {{ --ink:#1a1a1a; --muted:#6b7280; --line:#e5e7eb; --bg:#fff; }}
  * {{ box-sizing: border-box; }}
  body {{ margin:0; background:#fafafa; color:var(--ink);
         font: 15px/1.65 "Noto Sans SC", "Microsoft YaHei", system-ui, sans-serif; }}
  .wrap {{ max-width: 940px; margin: 0 auto; padding: 48px 24px 80px; }}
  h1 {{ font-size: 26px; margin: 0 0 6px; letter-spacing:-0.02em; }}
  .sub {{ color: var(--muted); font-size: 14px; margin-bottom: 36px; }}
  h2 {{ font-size: 15px; text-transform: uppercase; letter-spacing:.08em;
        color: var(--muted); margin: 44px 0 14px; font-weight:600; }}
  .kpis {{ display:grid; grid-template-columns: repeat(4,1fr); gap:14px; }}
  .kpi {{ background:var(--bg); border:1px solid var(--line); border-radius:10px;
          padding:16px 18px; }}
  .kpi .k {{ font-size:12px; color:var(--muted); }}
  .kpi .v {{ font-size:24px; font-weight:650; letter-spacing:-0.02em; margin-top:2px; }}
  .card {{ background:var(--bg); border:1px solid var(--line); border-radius:10px;
           padding:20px 22px; }}
  table {{ width:100%; border-collapse:collapse; font-size:13.5px; }}
  th, td {{ text-align:left; padding:8px 10px; border-bottom:1px solid var(--line); }}
  th {{ color:var(--muted); font-weight:600; font-size:12px; text-transform:uppercase;
        letter-spacing:.05em; }}
  td.num {{ text-align:right; font-variant-numeric: tabular-nums; }}
  .dot {{ display:inline-block; width:8px; height:8px; border-radius:50%;
          margin-right:7px; vertical-align:1px; }}
  .legend {{ display:flex; gap:18px; flex-wrap:wrap; margin-bottom:14px;
             font-size:13px; color:var(--muted); }}
  .legend i {{ display:inline-block; width:10px; height:10px; border-radius:2px;
              margin-right:6px; vertical-align:0; }}
  .note {{ font-size:13px; color:var(--muted); margin-top:12px; }}
  code {{ background:#f3f4f6; padding:1px 5px; border-radius:4px; font-size:13px; }}
</style>
</head>
<body>
<div class="wrap">
  <h1>输出轴案例：心 / 锁暝 / 守岸人</h1>
  <div class="sub">25 秒轴的逐动作伤害归因 · 数据来自手拉配装表，倍率已与官方数据源核对</div>

  <div class="kpis">
    <div class="kpi"><div class="k">团队总伤</div><div class="v">{total}</div></div>
    <div class="kpi"><div class="k">25 秒 DPS</div><div class="v">{dps}</div></div>
    <div class="kpi"><div class="k">动作数</div><div class="v">{n_actions}</div></div>
    <div class="kpi"><div class="k">单动作最高占比</div><div class="v">{top_share}</div></div>
  </div>

  <h2>三人在总伤中的占比</h2>
  <div class="card">{owner_svg}</div>

  <h2>每个动作的伤害贡献</h2>
  <div class="card">{action_legend}{bar_svg}</div>

  <h2>明细</h2>
  <div class="card">
    <table>
      <thead><tr><th>动作</th><th>归属</th><th>类型</th>
        <th style="text-align:right">伤害</th><th style="text-align:right">占比</th>
        <th style="text-align:right">取列</th></tr></thead>
      <tbody>{rows}</tbody>
    </table>
    <div class="note">「取列」标注该动作取的是期望列还是必暴击列 ——
      必暴击角色用后者，两者不能混用。</div>
  </div>

  <h2>模型说明</h2>
  <div class="card">
    <p style="margin:0 0 10px">伤害按九个乘区连乘：</p>
    <p style="margin:0 0 12px"><code>倍率 × 攻击 × 暴击 × (1+增伤) × (1+加深) × 防御 × 抗性 × 终伤 × (1+易伤)</code></p>
    <p style="margin:0">其中防御区为
      <code>(己方等级+100) / ((己方等级+100) + (敌方等级+99))</code>，
      敌人表里 90 级敌人的防御是 <code>1512 = 8×189</code>，
      所以偏移是 <b>+100 / +99</b> 而不是对称的 +100 / +100。</p>
    <div class="note">{check}</div>
  </div>
</div>
</body>
</html>
"""


def main():
    xlsx = Path(sys.argv[1] if len(sys.argv) > 1 else r"data\user_builds.xlsx")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "docs/demo/index.html")

    rows, totals = load(xlsx)
    actions, per_owner, total, check = build(rows, totals)

    table_rows = []
    for a in actions:
        table_rows.append(
            f'<tr><td>{a["label"]}</td>'
            f'<td><span class="dot" style="background:{a["color"]}"></span>{a["owner"]}</td>'
            f'<td>{a["type"]}</td>'
            f'<td class="num">{a["damage"]:,.0f}</td>'
            f'<td class="num">{a["damage"]/total*100:.2f}%</td>'
            f'<td class="num">{"期望" if a["column"]=="E" else "必暴击"}</td></tr>')

    action_legend = '<div class="legend">' + "".join(
        '<span><i style="background:%s"></i>%s</span>' % (GROUPS[o]["color"], o)
        for o in GROUPS
    ) + '</div>'

    check_txt = ""
    if check:
        parts = [f'{c["owner"]} 自算 {fmt(c["mine"])} vs 表内 {fmt(c["sheet"])}'
                 f'（差 {c["diff"]:+.2f}%）' for c in check]
        check_txt = "与手拉表自带的总伤公式交叉验证：" + "；".join(parts)

    html = HTML.format(
        total=fmt(total), dps=fmt(total / FIGHT_SECONDS),
        n_actions=len(actions),
        top_share=f'{actions[0]["damage"]/total*100:.1f}%',
        owner_svg=owner_svg(per_owner, total),
        bar_svg=bar_svg(actions, total),
        rows="\n      ".join(table_rows),
        check=check_txt,
        action_legend=action_legend,
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    io.open(out, "w", encoding="utf-8").write(html)
    print(f"wrote {out}  ({len(html):,} chars)")
    print(f"total={total:,.0f}  dps={total/FIGHT_SECONDS:,.0f}  actions={len(actions)}")
    for c in check:
        print(f"  check {c['owner']}: {c['mine']:,.0f} vs {c['sheet']:,.0f}  {c['diff']:+.3f}%")


if __name__ == "__main__":
    main()

