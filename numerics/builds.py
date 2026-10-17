"""Build enumeration and comparison tables ("拉表").

A damage calculator answers "how much damage does this build do". A build
table answers the question a player actually has: *"out of everything I
could put together, which one is best, and by how much"*.

Three things separate a useful table from a wall of numbers:

1. **Ranking with a relative baseline** -- every row shows ``% of best``,
   because "+12% over the next best" is a decision and "+18,432 damage"
   is not.
2. **A per-row diagnosis** -- each row carries the zone that is currently
   the best place to spend the next point, so the table tells you what to
   change, not just what is.
3. **Written conclusions** -- :func:`insights` turns the table back into
   sentences, because that is what ends up in a report.
"""
from __future__ import annotations

import csv
import itertools
from dataclasses import dataclass, field
from math import prod
from pathlib import Path
from typing import Iterable, Mapping, Optional, Sequence

from .damage import DamageModel, Zone


def crit_multiplier(crit_rate: float, crit_damage: float) -> float:
    """Expected crit zone value: ``1 + crit_rate * crit_damage``.

    ``crit_damage=1.10`` means a crit hits for 210% of normal damage.
    """
    if not 0.0 <= crit_rate <= 1.0:
        raise ValueError("crit_rate must be in [0, 1]")
    if crit_damage < 0:
        raise ValueError("crit_damage must be >= 0")
    return 1.0 + crit_rate * crit_damage


@dataclass(frozen=True)
class Build:
    """One candidate configuration: absolute zone values keyed by zone name."""

    name: str
    overrides: Mapping[str, float] = field(default_factory=dict)

    def __post_init__(self) -> None:
        for name, value in self.overrides.items():
            if value <= 0:
                raise ValueError(f"build {self.name!r}: zone {name!r} must be > 0")


def apply_build(model: DamageModel, build: Build) -> DamageModel:
    """Return a copy of ``model`` with the build's overrides applied."""
    known = {z.name for z in model.zones}
    unknown = set(build.overrides) - known
    if unknown:
        raise KeyError(f"build {build.name!r} references unknown zones: {sorted(unknown)}")
    zones = [
        Zone(z.name, build.overrides.get(z.name, z.value), z.per_point_gain, z.note)
        for z in model.zones
    ]
    return DamageModel(f"{model.label} / {build.name}".strip(" /"), zones)


def enumerate_builds(grid: Mapping[str, Sequence[float]],
                     max_builds: int = 20000) -> list[Build]:
    """Cartesian product of ``{zone: [candidate values]}``.

    Deterministic order (zones sorted by name) so runs are reproducible.
    """
    if not grid:
        raise ValueError("grid must not be empty")
    names = sorted(grid)
    options = [list(grid[n]) for n in names]
    for n, opts in zip(names, options):
        if not opts:
            raise ValueError(f"zone {n!r} has no candidate values")
    total = prod(len(o) for o in options)
    if total > max_builds:
        raise ValueError(
            f"grid would produce {total:,} builds, above max_builds={max_builds:,}; "
            "narrow the grid or raise the limit"
        )
    builds: list[Build] = []
    for combo in itertools.product(*options):
        name = ", ".join(f"{n}={v:g}" for n, v in zip(names, combo))
        builds.append(Build(name, dict(zip(names, combo))))
    return builds


def build_table(model: DamageModel, builds: Iterable[Build]) -> list[dict]:
    """Evaluate every build and rank the results.

    Each row carries the damage, the zone values, the relative performance
    against the best build, and the best place to spend the next point.
    """
    rows: list[dict] = []
    for build in builds:
        m = apply_build(model, build)
        rows.append({
            "build": build.name,
            "damage": m.total(),
            "zones": m.values(),
            "best_next_zone": m.cheapest_zone(),
            "marginal": dict(m.marginal_table()),
        })
    if not rows:
        raise ValueError("no builds to evaluate")
    rows.sort(key=lambda r: r["damage"], reverse=True)
    best = rows[0]["damage"]
    for rank, row in enumerate(rows, start=1):
        row["rank"] = rank
        row["pct_of_best"] = row["damage"] / best
        row["gap_to_best"] = row["damage"] / best - 1.0
    return rows


# ------------------------------------------------------------------ exports
def _columns(zone_order: Sequence[str]) -> list[str]:
    return ["rank", "build", *zone_order, "damage", "pct_of_best", "best_next_zone"]


def to_csv(rows: Sequence[dict], path: str | Path,
           zone_order: Optional[Sequence[str]] = None) -> Path:
    """Write the table as CSV (UTF-8 with BOM so Excel opens it correctly)."""
    if zone_order is None:
        zone_order = list(rows[0]["zones"])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(_columns(zone_order))
        for row in rows:
            writer.writerow([row["rank"], row["build"],
                             *[f"{row['zones'][z]:.4f}" for z in zone_order],
                             f"{row['damage']:.2f}",
                             f"{row['pct_of_best']:.6f}",
                             row["best_next_zone"]])
    return path


def to_markdown(rows: Sequence[dict], path: str | Path,
                zone_order: Optional[Sequence[str]] = None,
                top: int = 15) -> Path:
    """Write the top ``top`` rows as a Markdown table."""
    if zone_order is None:
        zone_order = list(rows[0]["zones"])
    head = ["Rank", "Build", *zone_order, "Damage", "% of best", "Best next zone"]
    lines = ["| " + " | ".join(head) + " |",
             "|" + "|".join(["---"] * len(head)) + "|"]
    for row in rows[:top]:
        cells = [str(row["rank"]), row["build"].replace("|", "\\|")]
        cells += [f"{row['zones'][z]:.3f}" for z in zone_order]
        cells += [f"{row['damage']:,.1f}", f"{row['pct_of_best'] * 100:.2f}%",
                  row["best_next_zone"]]
        lines.append("| " + " | ".join(cells) + " |")
    if len(rows) > top:
        lines.append("")
        lines.append(f"*{len(rows)} builds total; showing the top {top}.*")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def to_xlsx(rows: Sequence[dict], path: str | Path,
            zone_order: Optional[Sequence[str]] = None) -> Path:
    """Write the table as a formatted workbook.

    ``openpyxl`` is an optional extra -- the toolkit itself does not need it.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for XLSX export: pip install openpyxl") from exc

    if zone_order is None:
        zone_order = list(rows[0]["zones"])
    header = ["rank", "build", *zone_order, "damage", "pct_of_best", "best_next_zone"]
    wb = Workbook()
    ws = wb.active
    ws.title = "builds"

    accent = "1F4E79"
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    for j, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font = Font(bold=True, color=accent)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")

    for i, row in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=row["rank"])
        ws.cell(row=i, column=2, value=row["build"])
        for k, z in enumerate(zone_order):
            c = ws.cell(row=i, column=3 + k, value=round(row["zones"][z], 4))
            c.number_format = "0.0000"
        col = 3 + len(zone_order)
        c = ws.cell(row=i, column=col, value=round(row["damage"], 2))
        c.number_format = "#,##0.0"
        c = ws.cell(row=i, column=col + 1, value=round(row["pct_of_best"], 6))
        c.number_format = "0.00%"
        ws.cell(row=i, column=col + 2, value=row["best_next_zone"])

    ws.freeze_panes = "C2"
    for j, name in enumerate(header, start=1):
        width = 46 if name == "build" else (16 if name in zone_order else 13)
        ws.column_dimensions[get_column_letter(j)].width = width

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ----------------------------------------------------------------- insights
def insights(rows: Sequence[dict], base_model: Optional[DamageModel] = None) -> dict:
    """Turn a ranked table into the sentences that belong in a report."""
    if not rows:
        raise ValueError("no rows")
    best, runner = rows[0], (rows[1] if len(rows) > 1 else None)
    worst = rows[-1]

    zone_counts: dict[str, int] = {}
    for row in rows:
        z = row["best_next_zone"]
        zone_counts[z] = zone_counts.get(z, 0) + 1
    common_zone = max(zone_counts, key=zone_counts.get)

    out = {
        "best_build": best["build"],
        "best_damage": best["damage"],
        "worst_build": worst["build"],
        "worst_damage": worst["damage"],
        "spread": best["damage"] / worst["damage"] - 1.0,
        "n_builds": len(rows),
        "most_common_next_zone": common_zone,
        "share_needing_that_zone": zone_counts[common_zone] / len(rows),
        "best_next_zone_of_best": best["best_next_zone"],
    }
    if runner is not None:
        out["runner_up_build"] = runner["build"]
        out["gap_over_runner_up"] = best["damage"] / runner["damage"] - 1.0

    if base_model is not None:
        base = base_model.total()
        out["baseline_damage"] = base
        out["best_over_baseline"] = best["damage"] / base - 1.0
    return out


def format_insights(info: Mapping) -> list[str]:
    """Render :func:`insights` output as short human-readable lines."""
    lines = [
        f"{info['n_builds']} builds evaluated; best and worst differ by "
        f"{info['spread'] * 100:.1f}%.",
        f"Best build: {info['best_build']} ({info['best_damage']:,.1f})"
        + (f", {info['gap_over_runner_up'] * 100:.2f}% ahead of the runner-up."
           if "gap_over_runner_up" in info else "."),
        f"In {info['share_needing_that_zone'] * 100:.0f}% of builds the best next "
        f"investment is the same zone: {info['most_common_next_zone']}.",
    ]
    if "best_over_baseline" in info:
        lines.append(
            f"The best build is {info['best_over_baseline'] * 100:+.2f}% over the "
            f"baseline configuration."
        )
    return lines

