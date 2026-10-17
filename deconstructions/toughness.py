# -*- coding: utf-8 -*-
"""《鸣潮》敌人拆解 14：韧性档位与恢复。

第 03 篇拆角色削韧，这篇拆敌人韧性，两侧合起来才能算破韧成本。

用法:
    python deconstructions/toughness.py <数据表.xlsx> --json out.json
"""
from __future__ import annotations

import io
import json
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

SHEET = "敌对属性列表"
C_LV, C_TOUGH, C_RECOVER = 3, 17, 18
SENTINEL = 9999

# 第 03 篇测得的中位削韧
POISE = {"大招": 400, "技能": 40, "普攻": 26.7}


def load(xlsx):
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    ws = wb[SHEET]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or len(r) <= C_RECOVER:
            continue
        lv, t, rec = r[C_LV], r[C_TOUGH], r[C_RECOVER]
        if not isinstance(t, (int, float)) or t <= 0:
            continue
        out.append({"level": lv if isinstance(lv, (int, float)) else None,
                    "toughness": t,
                    "recover": rec if isinstance(rec, (int, float)) else None})
    wb.close()
    return out


def analyse(rows):
    tough = Counter(r["toughness"] for r in rows)
    recover = Counter(r["recover"] for r in rows if r["recover"] is not None)

    by_level = defaultdict(list)
    for r in rows:
        if r["level"]:
            by_level[int(r["level"])].append(r["toughness"])

    levels = []
    for lv in sorted(by_level):
        if lv in (40, 60, 80, 90, 100, 110, 120):
            v = sorted(by_level[lv])
            levels.append({"level": lv, "count": len(v),
                           "median": v[len(v) // 2]})

    # 破韧成本
    tiers = [t for t, n in tough.most_common() if t < SENTINEL][:6]
    cost = []
    for t in sorted(tiers):
        cost.append({"toughness": t,
                     "ults": round(t / POISE["大招"], 1),
                     "skills": round(t / POISE["技能"], 1),
                     "normals": round(t / POISE["普攻"], 1)})

    return {
        "rows": len(rows),
        "toughness_values": dict(tough.most_common()),
        "recover_values": dict(recover.most_common()),
        "by_level": levels,
        "cost": cost,
        "sentinel": SENTINEL,
    }


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else r"data\\鸣潮动作数据汇总.xlsx"
    res = analyse(load(xlsx))

    if "--json" in sys.argv:
        out = sys.argv[sys.argv.index("--json") + 1]
        io.open(out, "w", encoding="utf-8").write(
            json.dumps(res, ensure_ascii=False, indent=1))

    L = ["敌人 %d 条" % res["rows"], "", "== 韧性档位 =="]
    for v, n in res["toughness_values"].items():
        L.append("  %6g  %4d 条" % (v, n))
    L.append("")
    L.append("== 韧性恢复 ==")
    for v, n in res["recover_values"].items():
        L.append("  %6g  %4d 条" % (v, n))
    L.append("")
    L.append("韧性 | 大招 | 技能 | 普攻")
    for c in res["cost"]:
        L.append("%6g | %4.1f | %5.1f | %5.1f"
                 % (c["toughness"], c["ults"], c["skills"], c["normals"]))
    io.open("data/_toughness_summary.txt", "w", encoding="utf-8").write("\n".join(L))
    print("ok")


if __name__ == "__main__":
    main()
