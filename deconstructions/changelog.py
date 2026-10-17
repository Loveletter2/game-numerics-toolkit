# -*- coding: utf-8 -*-
"""《鸣潮》拆解 11：更新日志里改的是什么。

一条运营中的动作游戏，两年多里改过些什么、改的是谁、改的是哪一类。

用法:
    python deconstructions/changelog.py <数据表.xlsx> [角色名json] --json out.json
"""
from __future__ import annotations

import glob
import io
import json
import os
import sys
from collections import Counter

from openpyxl import load_workbook

SHEET = "更新变动"

# 改动内容的类型；命中即归类，顺序即优先级
KINDS = [
    ("弹刀/格挡", ("弹刀", "格挡")),
    ("无敌与闪避", ("无敌", "极限闪避")),
    ("敌人数值", ("生命值", "抗性", "暗抗", "物抗", "防御")),
    ("索敌/优先级", ("索敌", "优先级")),
    ("判定与子弹", ("子弹", "判定")),
    ("位置与偏移", ("偏移", "位置")),
    ("时间与帧", ("帧", "时长", "时间", "冷却")),
    ("削韧与韧性", ("削韧", "韧性", "白条")),
    ("伤害与倍率", ("伤害", "倍率")),
    ("范围与大小", ("大小", "范围", "半径", "缩放")),
]


def load_changes(xlsx):
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    ws = wb[SHEET]
    out = []
    date = version = None
    for r in ws.iter_rows(min_row=1, values_only=True):
        if not r or all(x is None for x in r):
            continue
        if r[0]:
            date = str(r[0]).strip()
        if len(r) > 2 and r[2]:
            version = str(r[2]).strip()
        if not r[1]:
            continue
        out.append({"date": date, "version": version, "text": str(r[1]).strip()})
    wb.close()
    return out


def load_names(dirpath):
    names = set()
    for fp in glob.glob(os.path.join(dirpath, "*.json")):
        try:
            d = json.load(io.open(fp, encoding="utf-8"))
        except Exception:
            continue
        n = d.get("name")
        if isinstance(n, str) and n:
            names.add(n)
        zh = d.get("zh")
        if isinstance(zh, dict) and zh.get("name"):
            names.add(zh["name"])
    return {n for n in names if len(n) >= 2}


def kind_of(text):
    for label, keys in KINDS:
        if any(k in text for k in keys):
            return label
    return "其他"


def analyse(rows, names):
    per_version = Counter((r["date"], r["version"]) for r in rows)
    versions = sorted(per_version)

    chars = Counter()
    for r in rows:
        hit = [n for n in names if n in r["text"]]
        if hit:
            # 取最长匹配，避免「长离」被「离」之类的短名误伤
            chars[max(hit, key=len)] += 1

    kinds = Counter(kind_of(r["text"]) for r in rows)
    by_year = Counter((r["date"] or "?").split(".")[0] for r in rows)

    sizes = sorted(per_version.values())
    return {
        "rows": len(rows),
        "versions": len(versions),
        "first": versions[0] if versions else None,
        "last": versions[-1] if versions else None,
        "median_per_version": sizes[len(sizes) // 2] if sizes else None,
        "max_per_version": sizes[-1] if sizes else None,
        "top_chars": chars.most_common(10),
        "kinds": dict(kinds.most_common()),
        "by_year": dict(sorted(by_year.items())),
    }


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else r"data\鸣潮动作数据汇总.xlsx"
    nandir = sys.argv[2] if len(sys.argv) > 2 else r"data\nanoka\3.7\zh_character"
    res = analyse(load_changes(xlsx), load_names(nandir))

    if "--json" in sys.argv:
        out = sys.argv[sys.argv.index("--json") + 1]
        io.open(out, "w", encoding="utf-8").write(
            json.dumps(res, ensure_ascii=False, indent=1))

    L = ["改动条目 %d 条，覆盖 %d 次更新" % (res["rows"], res["versions"]),
         "单次更新条目数 中位 %s，最多 %s"
         % (res["median_per_version"], res["max_per_version"]),
         "", "== 按年份 =="]
    for y, n in res["by_year"].items():
        L.append("  20%s 年：%d 条" % (y, n))
    L.append("")
    L.append("== 改动类型 ==")
    for k, n in res["kinds"].items():
        L.append("  %-14s %4d 条  %4.1f%%"
                 % (k, n, n / res["rows"] * 100))
    L.append("")
    L.append("== 被改最多的角色 ==")
    for name, n in res["top_chars"]:
        L.append("  %-10s %3d 条" % (name, n))
    io.open("data/_changelog_summary.txt", "w", encoding="utf-8").write("\n".join(L))
    print("ok")


if __name__ == "__main__":
    main()
