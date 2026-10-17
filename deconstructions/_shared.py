# -*- coding: utf-8 -*-
"""动作数据表的共用读取器。

表是社区整理的官方动作数据，每秒 60 帧。所有拆解脚本都从这里取数，
免得每个文件各写一遍列索引。

用法:
    from _shared import load_actions, FPS
    rows = load_actions("path/to/动作数据表.xlsx")
"""
from __future__ import annotations

from openpyxl import load_workbook

FPS = 60

# 列索引（0 起）。表头见 docs 或 README。
COLS = {
    "group": 0,          # 名称（体型骨架，不是角色名）
    "id": 1,
    "action": 2,
    "startup": 4,        # 发生帧
    "duration": 5,       # 持续帧
    "hitstop_self": 7,
    "iframe_start": 10,
    "iframe_dur": 11,
    "end": 15,           # 动作结束帧
    "poise": 19,         # 削韧值
    "ult": 21,           # 大招回收
    "concerto": 22,      # 协奏回收
    "core1": 23,
    "core2": 24,
    "core3": 25,
}


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return None
    return None


def load_actions(xlsx: str) -> list[dict]:
    """把「角色-女」「角色-男」两页合成一张扁平的动作表。"""
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    rows: list[dict] = []
    for sheet in ("角色-女", "角色-男"):
        if sheet not in wb.sheetnames:
            continue
        for r in wb[sheet].iter_rows(min_row=2, values_only=True):
            if not r or len(r) <= COLS["core3"]:
                continue
            action = r[COLS["action"]]
            if not action or not str(action).strip():
                continue
            item = {"sheet": sheet, "action": str(action).strip()}
            for key, idx in COLS.items():
                if key == "action":
                    continue
                v = r[idx]
                item[key] = v.strip() if isinstance(v, str) else v
            for key in ("startup", "duration", "iframe_start", "iframe_dur",
                        "end", "poise", "ult", "concerto", "core1", "core2",
                        "core3", "hitstop_self"):
                item[key] = _num(item.get(key))
            rows.append(item)
    wb.close()
    return rows


def seconds(frames: float | None) -> float | None:
    return None if frames is None else round(frames / FPS, 3)
