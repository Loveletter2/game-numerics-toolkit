# -*- coding: utf-8 -*-
"""《鸣潮》敌人拆解 06：血量成长与击杀时间。

用一张配装表的 DPS 去量敌人的血量，看看一套爆发能带走什么。

用法:
    python deconstructions/enemies.py <数据表.xlsx> [--dps 90695] --json out.json
"""
from __future__ import annotations

import io
import json
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

SHEET = "敌对属性列表"
C_LV, C_HP, C_DEF = 3, 4, 6
RES_COLS = {"物理": 7, "冰": 8, "火": 9, "雷": 10, "风": 11, "光": 12, "暗": 13}
KEY_LEVELS = (40, 60, 80, 90, 100, 110, 120)


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return None
    return None


def load(xlsx):
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    ws = wb[SHEET]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or len(r) <= max(RES_COLS.values()):
            continue
        lv, hp = num(r[C_LV]), num(r[C_HP])
        if not lv or not hp or hp <= 0:
            continue
        out.append({
            "level": int(lv),
            "hp": hp,
            "def": num(r[C_DEF]),
            "res": {k: num(r[i]) for k, i in RES_COLS.items()},
        })
    wb.close()
    return out


def tiers(rows, dps, min_share=3):
    """同一等级下被多个敌人共用的血量 —— 说明它是查表值而不是逐怪配置。"""
    from collections import Counter
    by_lv = {}
    for r in rows:
        by_lv.setdefault(r["level"], []).append(round(r["hp"]))
    out = []
    for lv in sorted(by_lv):
        if lv < 85:
            continue
        c = Counter(by_lv[lv])
        hp, n = c.most_common(1)[0]
        if n < min_share:
            continue
        out.append({"level": lv, "hp": hp, "shared_by": n,
                    "ttk": round(hp / dps, 1) if dps else None})
    return out


def analyse(rows, dps):
    by_lv = defaultdict(list)
    for r in rows:
        by_lv[r["level"]].append(r["hp"])

    curve = []
    for lv in sorted(by_lv):
        if lv not in KEY_LEVELS:
            continue
        vals = sorted(by_lv[lv])
        med = vals[len(vals) // 2]
        curve.append({
            "level": lv,
            "count": len(vals),
            "median_hp": round(med),
            "ttk_seconds": round(med / dps, 1) if dps else None,
        })

    res_hist = Counter()
    for r in rows:
        for v in r["res"].values():
            if v is None:
                continue
            res_hist[round(v, 2)] += 1

    return {
        "row_count": len(rows),
        "levels": sorted(by_lv),
        "curve": curve,
        "resistance_histogram": dict(sorted(res_hist.items())),
        "dps_assumption": dps,
        "tiers": tiers(rows, dps),
    }


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else r"data\鸣潮动作数据汇总.xlsx"
    dps = 90695.0
    if "--dps" in sys.argv:
        dps = float(sys.argv[sys.argv.index("--dps") + 1])
    res = analyse(load(xlsx), dps)

    if "--json" in sys.argv:
        out = sys.argv[sys.argv.index("--json") + 1]
        io.open(out, "w", encoding="utf-8").write(
            json.dumps(res, ensure_ascii=False, indent=1))

    L = ["敌人条目 %d 条，等级 %d～%d" % (res["row_count"],
                                          res["levels"][0], res["levels"][-1]),
         "假设队伍 DPS %.0f" % res["dps_assumption"], "",
         "== 血量成长与击杀时间 ==",
         "等级  条目  血量中位   TTK"]
    for c in res["curve"]:
        L.append("%4d %5d %10s %6ss" % (c["level"], c["count"],
                                        format(c["median_hp"], ","),
                                        c["ttk_seconds"]))
    L.append("")
    L.append("== 抗性取值分布（七种元素合计）==")
    for v, n in res["resistance_histogram"].items():
        L.append("  %5.2f  %5d 次" % (v, n))
    io.open("data/_enemies_summary.txt", "w", encoding="utf-8").write("\n".join(L))
    print("ok")


if __name__ == "__main__":
    main()
