# -*- coding: utf-8 -*-
"""《鸣潮》动作帧拆解：闪避类动作的无敌窗口。

数据源：社区整理的官方动作数据表（每秒 60 帧）。
脚本只读表，不修改任何数据。

用法:
    python deconstructions/frames.py <动作数据表.xlsx> --json out.json
"""
from __future__ import annotations

import io
import json
import sys
from collections import defaultdict

from openpyxl import load_workbook

FPS = 60
C_NAME, C_ACTION, C_IFRAME = 0, 2, 11
C_POISE, C_CONCERTO = 19, 22


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return None
    return None


def load(xlsx: str):
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    rows = []
    for sheet in ("角色-女", "角色-男"):
        if sheet not in wb.sheetnames:
            continue
        for r in wb[sheet].iter_rows(min_row=2, values_only=True):
            if not r or len(r) <= C_IFRAME:
                continue
            action = r[C_ACTION]
            if not action or not str(action).strip():
                continue
            rows.append({
                "group": str(r[C_NAME]).strip() if r[C_NAME] else None,
                "action": str(action).strip(),
                "iframe": num(r[C_IFRAME]),
                "poise": num(r[C_POISE]),
                "concerto": num(r[C_CONCERTO]),
            })
    wb.close()
    return rows


def dodge_windows(rows):
    """按动作名聚合「闪避」类动作的无敌帧。"""
    buckets = defaultdict(list)
    for r in rows:
        if "闪避" in r["action"] and r["iframe"]:
            buckets[r["action"]].append(r["iframe"])

    out = []
    for action, vals in buckets.items():
        out.append({
            "action": action,
            "frames": max(vals),
            "seconds": round(max(vals) / FPS, 3),
            "samples": len(vals),
            "consistent": len(set(vals)) == 1,
        })
    out.sort(key=lambda x: (-x["frames"], -x["samples"]))
    return out


def histogram(windows):
    """无敌窗口长度分布。"""
    h = defaultdict(int)
    for w in windows:
        h[w["frames"]] += w["samples"]
    return dict(sorted(h.items(), reverse=True))


def top_values(rows, key, limit=8, ceiling=1000):
    vals = [r for r in rows if r[key] and 0 < r[key] <= ceiling]
    vals.sort(key=lambda r: -r[key])
    return [{"action": r["action"], key: r[key]} for r in vals[:limit]]


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else r"data\鸣潮动作数据汇总.xlsx"
    rows = load(xlsx)
    windows = dodge_windows(rows)
    res = {
        "row_count": len(rows),
        "actions_with_iframe": sum(1 for r in rows if r["iframe"]),
        "dodge_windows": windows,
        "histogram": histogram(windows),
        "poise_top": top_values(rows, "poise"),
        "concerto_top": top_values(rows, "concerto", ceiling=100),
    }

    if "--json" in sys.argv:
        path = sys.argv[sys.argv.index("--json") + 1]
        io.open(path, "w", encoding="utf-8").write(
            json.dumps(res, ensure_ascii=False, indent=1))

    lines = [
        "动作总数 %d，其中带无敌帧 %d" % (res["row_count"], res["actions_with_iframe"]),
        "",
        "== 闪避类无敌窗口 ==",
    ]
    for w in windows[:20]:
        lines.append("  %-20s %2dF %5.2fs  样本%2d  %s" % (
            w["action"], w["frames"], w["seconds"], w["samples"],
            "一致" if w["consistent"] else "不一致"))
    lines.append("")
    lines.append("== 分布 ==")
    for f, n in res["histogram"].items():
        lines.append("  %3dF  %d 条" % (f, n))
    io.open("data/_frames_summary.txt", "w", encoding="utf-8").write("\n".join(lines))
    print("ok")


if __name__ == "__main__":
    main()
